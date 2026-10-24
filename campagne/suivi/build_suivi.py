# Génère le classeur de suivi RPVD (importé dans Google Sheets). Les formules sont dans l'onglet « Synthèse ».
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = Workbook()
HEAD = PatternFill('solid', fgColor='1F3864')
HF = Font(name='Arial', bold=True, color='FFFFFF')
INPUT = Font(name='Arial', color='0000FF')
BASE = Font(name='Arial')


def sheet(ws, headers, widths, rows=(), freeze='A2'):
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEAD
        c.font = HF
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 42
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = INPUT if c.value is not None and not str(c.value).startswith('=') else BASE


# 1. Guide
g = wb.active
g.title = 'Guide'
lines = [
    ('Suivi des publications RPVD', True),
    ("Une ligne par publication et par plateforme dans l'onglet Publications. Les onglets Synthèse et Entonnoir se calculent seuls.", False),
    ('Texte bleu : valeur saisie (statistiques lues sur la plateforme). Texte noir : formule.', False),
    ('', False),
    ('Mesures : lire les statistiques 24 h après la publication (colonne « Mesuré à » = 24 h), puis à 7 jours (7 j). Une ligne par mesure.', False),
    ("Ressenti : note de 1 à 5 donnée par l'équipe (1 = ne fonctionne pas, 5 = très bon signal). Observations : une phrase factuelle.", False),
    ('Source des données : pages de statistiques de chaque plateforme, lues le jour indiqué dans « Mesuré le ». Les lignes du Jour 1 viennent de Meta Business Suite le 16 octobre 2026.', False),
    ("Les statistiques de TikTok et Snapchat sont à compléter dès qu'elles sont disponibles.", False),
]
for i, (t, b) in enumerate(lines, 1):
    g.cell(i, 1, t).font = Font(name='Arial', bold=b, size=14 if b else 10)
g.column_dimensions['A'].width = 130

# 2. Publications
pub_headers = [
    'Date de publication', 'Jour n°', 'Plateforme', 'Format', 'Compte ou groupe', 'Heure', 'Accroche (3 premières secondes)',
    'Produit mis en avant', 'Lien de la publication', 'Mesuré le', 'Mesuré à', 'Vues', 'Spectateurs uniques', "Réactions (j'aime)",
    'Commentaires', 'Partages', 'Enregistrements', 'Clics sur le lien', 'Durée moyenne de visionnage (s)', '% moyen de visionnage',
    "Seconde d'abandon principale", 'Abonnés gagnés', "Tranche d'âge principale", 'Part de cette tranche', 'Ressenti (1 à 5)', 'Observations',
]
E = ''
pub_rows = [
    ['2026-10-15', 1, 'Facebook', 'Reel', 'RPVD Success', '21:47', 'Dimanche, 21 h.', 'Examen', E, '2026-10-16', '24 h', 227, 207, 0, 0, 0, E, 1, 3, 0.10, 4, 0, '45 à 54 ans', 0.315, 2,
     'Abandon dès 4 s : accroche trop lente. Public 100 % non-abonnés, 56 % entre 35 et 54 ans.'],
    ['2026-10-15', 1, 'Instagram', 'Reel', 'rpvdsuccess', '21:47', 'Dimanche, 21 h.', 'Examen', E, '2026-10-16', '24 h', 36, E, E, E, E, E, E, E, E, E, E, E, E, E,
     'Vues Instagram du même reel (publication croisée). Statistiques détaillées à relever.'],
    ['2026-10-15', 1, 'TikTok', 'Vidéo', 'RPVD', '21:50', 'Dimanche, 21 h.', 'Examen', E, '2026-10-16', '24 h', E, E, E, E, E, E, E, E, E, E, E, E, E, E,
     "En cours d'examen par TikTok au moment de la publication. À relever."],
    ['2026-10-15', 1, 'Facebook', 'Story', 'RPVD Success', '21:55', 'Dimanche, 21 h.', 'Examen', E, E, E, E, E, E, E, E, E, E, E, E, E, E, E, E, E,
     'Statistiques de story à relever.'],
    ['2026-10-15', 1, 'Facebook', 'Post de groupe', 'Parent entraide Sherbrooke', '21:50', 'Question aux parents', E, E, '2026-10-16', '24 h', 0, E, E, E, E, E, E, E, E, E, E, E, E, E,
     'Post supprimé par Facebook ou par le groupe presque aussitôt.'],
    ['2026-10-15', 1, 'Facebook', 'Post de groupe', 'Soutien scolaire et aide aux devoirs', '22:05', 'Question aux parents', E, E, E, E, E, E, E, E, E, E, E, E, E, E, E, E, E, E,
     "En attente d'approbation d'un administrateur."],
]
ws = wb.create_sheet('Publications')
sheet(ws, pub_headers, [14, 8, 14, 14, 28, 8, 32, 16, 30, 12, 10, 10, 12, 12, 12, 10, 12, 12, 14, 12, 14, 10, 16, 12, 10, 60], pub_rows, 'D2')
for r in range(2, 2000):
    ws.cell(r, 20).number_format = '0%'
    ws.cell(r, 24).number_format = '0.0%'
for rng, formula in [('C2:C2000', '"Facebook,Instagram,TikTok,Snapchat"'), ('D2:D2000', '"Reel,Story,Vidéo,Post de groupe,Carrousel"'), ('K2:K2000', '"24 h,7 j"')]:
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)

# 3. Synthèse (formules)
s = wb.create_sheet('Synthèse')
s.append(['Plateforme', 'Publications mesurées', 'Vues (total)', 'Spectateurs uniques', 'Réactions', 'Commentaires', 'Partages', 'Clics sur le lien',
          'Durée moyenne de visionnage (s)', 'Ressenti moyen (1 à 5)', "Taux d'interaction", 'Taux de clic'])
for c in s[1]:
    c.fill = HEAD
    c.font = HF
    c.alignment = Alignment(wrap_text=True, vertical='center')
for i, p in enumerate(['Facebook', 'Instagram', 'TikTok', 'Snapchat'], 2):
    s.cell(i, 1, p).font = BASE
    s.cell(i, 2, f'=COUNTIFS(Publications!$C:$C,$A{i},Publications!$L:$L,">0")')
    for col, src in [(3, 'L'), (4, 'M'), (5, 'N'), (6, 'O'), (7, 'P'), (8, 'R')]:
        s.cell(i, col, f'=SUMIFS(Publications!${src}:${src},Publications!$C:$C,$A{i})')
    s.cell(i, 9, f'=IFERROR(AVERAGEIFS(Publications!$S:$S,Publications!$C:$C,$A{i}),"")')
    s.cell(i, 10, f'=IFERROR(AVERAGEIFS(Publications!$Y:$Y,Publications!$C:$C,$A{i}),"")')
    s.cell(i, 11, f'=IF(C{i}=0,0,(E{i}+F{i}+G{i})/C{i})')
    s.cell(i, 12, f'=IF(C{i}=0,0,H{i}/C{i})')
    s.cell(i, 11).number_format = '0.0%'
    s.cell(i, 12).number_format = '0.0%'
    s.cell(i, 9).number_format = '0.0'
    s.cell(i, 10).number_format = '0.0'
s.cell(6, 1, 'Total').font = Font(name='Arial', bold=True)
for col in range(2, 9):
    L = get_column_letter(col)
    s.cell(6, col, f'=SUM({L}2:{L}5)').font = Font(name='Arial', bold=True)
s.cell(6, 11, '=IF(C6=0,0,(E6+F6+G6)/C6)')
s.cell(6, 12, '=IF(C6=0,0,H6/C6)')
s.cell(6, 11).number_format = '0.0%'
s.cell(6, 12).number_format = '0.0%'
s.cell(8, 1, 'Les taux se calculent sur les vues. Le taux de clic compte les clics vers rpvdsuccess.com divisés par les vues.').font = Font(name='Arial', italic=True)
for i, w in enumerate([16, 14, 12, 14, 12, 14, 10, 14, 18, 16, 14, 12], 1):
    s.column_dimensions[get_column_letter(i)].width = w
s.row_dimensions[1].height = 42
s.freeze_panes = 'B2'

# 4. Entonnoir
e = wb.create_sheet('Entonnoir')
sheet(e, ['Date', 'Source (src=)', 'Visites du site', 'Sondages envoyés', 'Démos regardées', 'Achats', "Chiffre d'affaires ($)", 'Taux visite vers sondage', 'Taux sondage vers achat', 'Observations'],
      [14, 22, 14, 16, 16, 10, 18, 18, 18, 60],
      [['2026-10-15', 'ig_reel, fb_reel, tiktok', E, E, E, E, E, E, E,
        'Visites : tableau de bord Netlify ou Google Analytics. Sondages et achats : page /admin.']])
for r in range(2, 2000):
    e.cell(r, 8, f'=IF(C{r}="","",IF(C{r}=0,0,D{r}/C{r}))')
    e.cell(r, 9, f'=IF(D{r}="","",IF(D{r}=0,0,F{r}/D{r}))')
    e.cell(r, 8).number_format = '0.0%'
    e.cell(r, 9).number_format = '0.0%'
    e.cell(r, 7).number_format = '$#,##0.00'

# 5. Ajustements
a = wb.create_sheet('Ajustements')
sheet(a, ['Date', 'Jour n°', "Constat (chiffre à l'appui)", 'Hypothèse', 'Changement fait au prompt du jour', 'Résultat observé le lendemain', 'Garder ?'],
      [14, 8, 50, 40, 50, 40, 10],
      [['2026-10-16', 2, 'Reel Jour 1 : durée moyenne de visionnage 3 s, abandon principal à 4 s, 227 vues, 1 clic.',
        "L'accroche « Dimanche, 21 h. » ne retient pas : le sujet n'apparaît pas dans les 2 premières secondes.",
        "À décider : faire apparaître la question « T'as étudié pour jeudi ? » dès 0 à 2 s.", E, E]])

# 6. Groupes Facebook
f = wb.create_sheet('Groupes Facebook')
sheet(f, ['Date', 'Groupe', 'Membres', "Statut d'adhésion", 'Texte publié (version)', 'Lien en commentaire', 'Visible après 10 min', 'Commentaires', 'Réponses données', 'Visites (?src=fb_groupe)', 'Observations'],
      [14, 40, 10, 18, 22, 16, 22, 12, 14, 16, 50],
      [['2026-10-15', 'Parent entraide Sherbrooke', 4700, 'Membre', 'Jour 1 [B]', 'Oui', 'Non (supprimé)', E, E, E, 'Supprimé par Facebook ou par le groupe, cause inconnue.'],
       ['2026-10-15', 'Soutien scolaire et aide aux devoirs', 51600, 'Membre', 'Jour 1 [A]', 'Non', "En attente d'approbation", E, E, E, 'Lien à ajouter une fois le post approuvé.']])

wb.save('suivi-rpvd.xlsx')
print('ok')
